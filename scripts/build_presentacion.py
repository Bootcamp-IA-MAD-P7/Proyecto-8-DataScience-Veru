"""Genera la presentación (.pptx) del proyecto para la exposición en clase.

Uso:
    uv run python scripts/build_presentacion.py --out presentacion.pptx
Después convertir a PDF con LibreOffice:
    soffice --headless --convert-to pdf --outdir . presentacion.pptx

Estructura (Idea de negocio + Parte técnica), alineada con README y SDD.
Diseño: fondo oscuro (#0f172a) y texto claro, en línea con el frontend.
"""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent

FONDO = RGBColor(0x0F, 0x17, 0x2A)
ACENTO = RGBColor(0x38, 0xBD, 0xF8)
BLANCO = RGBColor(0xEA, 0xEA, 0xEA)
GRIS = RGBColor(0x94, 0xA3, 0xB8)
VERDE = RGBColor(0x4A, 0xDE, 0x80)


def set_slide_bg(slide, color):
    """Pinta el fondo de la diapositiva con un color sólido."""
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = color


def add_bullets(slide, items, left, top, width, height, size=16, color=BLANCO):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = "•  " + item
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
    return tb


def add_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = ACENTO
    return tb


def add_center(slide, text, top, size, color, bold=False):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def new_slide(prs):
    """Crea una diapositiva en blanco con fondo oscuro."""
    blank = prs.slide_layouts[6]  # layout "en blanco"
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, FONDO)
    return slide


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "presentacion.pptx"))
    args = ap.parse_args()

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # Portada
    s = new_slide(prs)
    add_center(s, "Predicción del riesgo de ictus", 1.4, 36, BLANCO, True)
    add_center(s, "Machine Learning · Bootcamp Data Science / AI Developer", 2.4, 18, GRIS)
    add_center(s, "Proyecto 8 · 2026", 3.0, 16, ACENTO)

    # Sección Idea de negocio
    s = new_slide(prs)
    add_center(s, "Primera parte — Idea de negocio", 1.7, 30, ACENTO, True)
    add_center(s, "Presentación al cliente", 2.7, 18, GRIS)

    # Propósito / problema / público
    s = new_slide(prs)
    add_title(s, "Propósito · Problema · Público objetivo")
    add_bullets(s, [
        "Propósito: modelo de ML que, a partir de datos clínicos de un paciente, devuelve la probabilidad de sufrir un ictus.",
        "Problema que soluciona: cribado temprano del riesgo de ictus, donde la clase positiva (ictus) es minoritaria (4.98 %).",
        "Público objetivo: personal sanitario, equipos de desarrollo y pacientes en campañas de prevención.",
    ], 0.8, 1.4, 8.4, 3.4, 16)

    # User stories
    s = new_slide(prs)
    add_title(s, "User stories")
    add_bullets(s, [
        "Como personal sanitario, quiero conocer la probabilidad de ictus de un paciente, para priorizar la atención.",
        "Como desarrollador/a, quiero invocar la predicción desde un servicio HTTP (API), para integrarla en una aplicación.",
        "Como paciente, quiero una interfaz sencilla que muestre claramente mi nivel de riesgo.",
    ], 0.8, 1.4, 8.4, 3.4, 16)

    # Tecnologías
    s = new_slide(prs)
    add_title(s, "Tecnologías utilizadas")
    add_bullets(s, [
        "EDA: Pandas, Seaborn, Matplotlib, missingno (notebooks 01–07).",
        "Modelos: scikit-learn, CatBoost, XGBoost, imbalanced-learn.",
        "Productivización: FastAPI, Uvicorn, Pydantic (API REST + frontend).",
        "Base de datos: PostgreSQL (Docker) / SQLite (local).",
        "Tooling: uv, Makefile, pytest, Git, Docker Compose.",
    ], 0.8, 1.4, 8.4, 3.4, 16)

    # Estructura del código
    s = new_slide(prs)
    add_title(s, "Estructura del código")
    add_bullets(s, [
        "BACKEND/ — API FastAPI (main.py), adaptador BD (db.py) y frontend (static/).",
        "scripts/ — entrenamiento (train_*), comparativa (compare) y CLI (predict_cli).",
        "models/ — modelo final (catboost_final.pkl) + comparativos + informe.",
        "tests/ — suite de tests pytest + informe.",
        "SDD/ — decisiones de diseño (D1–D13).",
        "dockerfile + docker-compose.yml — contenedor y orquestación.",
    ], 0.8, 1.3, 8.4, 3.7, 15)

    # Despliegue
    s = new_slide(prs)
    add_title(s, "Uso de servidores de despliegue")
    add_bullets(s, [
        "Planificado en Render (Web Service + PostgreSQL).",
        "Ejecución actual en local: make api → http://127.0.0.1:8000 (frontend + API).",
        "El modelo y la app están listos para contenerizarse y desplegarse.",
    ], 0.8, 1.4, 8.4, 3.0, 16)

    # Nube / contenedores
    s = new_slide(prs)
    add_title(s, "Herramientas de nube y contenedores")
    add_bullets(s, [
        "Nube (AWS/GCP/Azure/S3/Firebase): No se utiliza — todo se ejecuta en local/Postgres.",
        "Contenedores (Docker): Configurado — docker-compose con 2 servicios (db + api).",
        "Orquestación: docker-compose up -d --build (make docker-up).",
        "Kubernetes: No se utiliza.",
    ], 0.8, 1.4, 8.4, 3.2, 16)

    # Repositorio y CI
    s = new_slide(prs)
    add_title(s, "Repositorio y CI")
    add_bullets(s, [
        "Repo: GitHub — Bootcamp-IA-MAD-P7/Proyecto-8-DataScience-Veru.",
        "Ramas: feature → dev → main (EDA, models, informe, API, database, frontend, docker).",
        "Commits descriptivos en español (feat/, fix/, chore/, docs/).",
        "GitHub Actions: No configurado — verificación manual con make test (24 tests).",
    ], 0.8, 1.4, 8.4, 3.2, 16)

    # Userflow
    s = new_slide(prs)
    add_title(s, "Flujo de usuario (userflow)")
    add_bullets(s, [
        "Vía CLI: python scripts/predict_cli.py --age ... --gender ... → probabilidad + veredicto.",
        "Vía API: POST /predict (JSON) → {probabilidad_ictus, clase, riesgo}.",
        "Vía Web: GET / → formulario con 10 campos → gauge + veredicto + historial.",
        "Historial: GET /predictions (guardado en BD).",
    ], 0.8, 1.4, 8.4, 3.4, 16)

    # Sección Parte técnica
    s = new_slide(prs)
    add_center(s, "Segunda parte — Parte técnica", 1.7, 30, ACENTO, True)
    add_center(s, "Presentación al equipo de desarrollo", 2.7, 18, GRIS)

    # Implementaciones importantes
    s = new_slide(prs)
    add_title(s, "Implementaciones importantes")
    add_bullets(s, [
        "Pipeline de preprocesado + modelo (evita data leakage, D6).",
        "Desbalance: scale_pos_weight / class_weight / SMOTE (D2).",
        "Overfitting: |train−test| ≤ 5 puntos (D3.3).",
        "Validación cruzada StratifiedKFold(5) (D5.1) y GridSearchCV (D5.2).",
        "CLI + API REST + Frontend web + BD historial + Docker.",
    ], 0.8, 1.3, 8.4, 3.6, 15)

    # Base de datos
    s = new_slide(prs)
    add_title(s, "Diseño de la base de datos (D11)")
    add_bullets(s, [
        "Origen 1: data/stroke_dataset.csv (entrenamiento, solo lectura).",
        "Origen 2: tabla predictions en PostgreSQL (producción) o SQLite (local).",
        "Entidad única: prediction (id PK, 10 features + probabilidad, clase, riesgo, created_at).",
        "Sin relaciones ni claves foráneas (modelo sobre paciente puntual).",
        "Endpoints: POST /predict (guarda) y GET /predictions (historial).",
    ], 0.8, 1.3, 8.4, 3.6, 15)

    # Paradigma, estructura, patrones
    s = new_slide(prs)
    add_title(s, "Paradigma, estructura y patrones")
    add_bullets(s, [
        "Desarrollo funcional por script, OOP donde aporta (Pydantic, Pipeline).",
        "Patrón: Pipeline de scikit-learn (preprocesado + modelo).",
        "Bootstrap: BACKEND/main.py carga el modelo al arrancar (joblib.load).",
        "Config por constantes de módulo (THRESHOLD, RANDOM_STATE, FEATURES).",
    ], 0.8, 1.4, 8.4, 3.4, 16)

    # Convenciones / repo / docs
    s = new_slide(prs)
    add_title(s, "Convenciones, repositorio y documentación")
    add_bullets(s, [
        "Naming: scripts snake_case (train_, predict_, compare_), funciones verbos, clases UpperCamelCase.",
        "Modelos: make train-* → models/<modelo>.pkl; constantes en MAYÚSCULAS_SNAKE.",
        "Docs: README.md, SDD (D1–D13), models/INFORME_MODELOS.md, tests/INFORME_TESTS.md.",
        "API servidor: /health, /predict, /predictions + Swagger /docs.",
    ], 0.8, 1.3, 8.4, 3.6, 15)

    # Modelo final y cierre
    s = new_slide(prs)
    add_title(s, "Modelo final y cierre")
    add_bullets(s, [
        "Modelo: CatBoost regularizado (config 'med').",
        "Métricas en test: Recall 0.86 · Precision 0.146 · F1 0.250.",
        "Overfitting: |train − test| ≤ 1.4 puntos (límite 5). ✅",
        "Tests: 24 tests pasando (make test).",
        "Conclusiones: age domina la importancia; desbalance limita la precisión.",
        "Vías futuras: más variables, despliegue completo en Render.",
    ], 0.8, 1.3, 8.4, 3.7, 15)

    prs.save(args.out)
    print(f"Presentación guardada en {args.out}")


if __name__ == "__main__":
    main()
